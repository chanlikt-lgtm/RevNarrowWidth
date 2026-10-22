"""Build a slide-deck summary of the Li 1990 semi-recessed-oxide narrow-channel
conformal-mapping reproduction (companion to narrow_channel_semi_recessed_report.pdf)."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

R = r"E:\claude\RAG\reports"
OUT = f"{R}/li1990_conformal_mapping_summary.pptx"

NAVY = RGBColor(0x1A, 0x2A, 0x5E)
GREY = RGBColor(0x44, 0x44, 0x44)
DKRED = RGBColor(0xA0, 0x20, 0x20)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def add_slide():
    return prs.slides.add_slide(BLANK)


def title(slide, text, size=32, top=0.35, color=NAVY):
    tb = slide.shapes.add_textbox(Inches(0.5), Inches(top), Inches(12.3), Inches(1.0))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = True
    p.font.color.rgb = color
    return tb


def subtitle(slide, text, top=1.05, size=16, color=GREY, italic=True):
    tb = slide.shapes.add_textbox(Inches(0.5), Inches(top), Inches(12.3), Inches(0.5))
    p = tb.text_frame.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.italic = italic
    p.font.color.rgb = color
    return tb


def bullets(slide, items, left=0.6, top=1.7, width=12.0, height=5.3, size=16.5,
            bullet_char="• "):
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, (txt, lvl) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = (bullet_char if lvl == 0 else "   – ") + txt
        p.font.size = Pt(size if lvl == 0 else size - 1.5)
        p.font.color.rgb = RGBColor(0x20, 0x20, 0x20)
        p.space_after = Pt(8 if lvl == 0 else 4)
    return tb


def pic(slide, path, left, top, width=None, height=None):
    kw = {}
    if width:
        kw["width"] = Inches(width)
    if height:
        kw["height"] = Inches(height)
    return slide.shapes.add_picture(path, Inches(left), Inches(top), **kw)


def footer(slide, text="Li, Hong, Cheng & Chan (1990), IEEE TED 37(3), pp. 692-701"):
    tb = slide.shapes.add_textbox(Inches(0.5), Inches(7.05), Inches(12.3), Inches(0.35))
    p = tb.text_frame.paragraphs[0]
    p.text = text
    p.font.size = Pt(10)
    p.font.color.rgb = GREY


# ---------------------------------------------------------------- 1. Title
s = add_slide()
tb = s.shapes.add_textbox(Inches(0.8), Inches(2.3), Inches(11.7), Inches(1.3))
p = tb.text_frame.paragraphs[0]
p.text = "Narrow-Channel Effect via Conformal Mapping"
p.font.size = Pt(40)
p.font.bold = True
p.font.color.rgb = NAVY
p.alignment = PP_ALIGN.CENTER
tb2 = s.shapes.add_textbox(Inches(0.8), Inches(3.5), Inches(11.7), Inches(0.9))
p2 = tb2.text_frame.paragraphs[0]
p2.text = "Reproduction & finite-difference verification of Li, Hong, Cheng & Chan (1990),\nsemi-recessed field-oxide MOSFETs"
p2.font.size = Pt(18)
p2.font.color.rgb = GREY
p2.alignment = PP_ALIGN.CENTER
tb3 = s.shapes.add_textbox(Inches(0.8), Inches(6.4), Inches(11.7), Inches(0.6))
p3 = tb3.text_frame.paragraphs[0]
p3.text = "scripts/narrow_semi_recessed/  —  E:\\claude\\RAG"
p3.font.size = Pt(13)
p3.font.color.rgb = GREY
p3.alignment = PP_ALIGN.CENTER

# ---------------------------------------------------------------- 2. Problem / geometry
s = add_slide()
title(s, "The Physical Problem")
subtitle(s, "Semi-recessed field-oxide isolation — narrow-channel MOSFETs")
bullets(s, [
    ("Field oxide of thickness tᶠ sits beside the gate; recess depth t₁ controls how far the oxide extends down the sidewall.", 0),
    ("Fully-recessed (t₁ = tᶠ−tᵍ) and non-recessed (t₁→0) structures are the SAME model at limiting t₁ — no separate formula needed.", 0),
    ("Channel depletion depth t_d vs. recess depth t₁ sets the regime:", 0),
    ("Case I (t_d < t₁): depletion stays inside the sidewall → inverse narrow-width effect (ΔVth < 0)", 1),
    ("Case II (t_d > t₁): depletion spills laterally under the field oxide, a quarter-circle “wing” of radius W_c = t_d−t₁ → conventional effect (ΔVth > 0)", 1),
    ("Goal: solve ∇²φ = 0 in the oxide cross-section to get the sidewall/lateral flux that sets Vth(W).", 0),
], top=1.75, size=17)
footer(s)

# ---------------------------------------------------------------- 3. Conformal mapping method
s = add_slide()
title(s, "Why Conformal Mapping")
subtitle(s, "Turning a messy 2-D corner geometry into a solvable strip problem")
bullets(s, [
    ("Laplace's equation ∇²φ=0 is solved on a semi-infinite half-strip (the field-oxide cross-section) with mixed Dirichlet boundary data on 4 edges.", 0),
    ("The map u+iv = cos[π(ix−y)/tᶠ] sends this half-strip onto the upper half-plane, where the same BVP becomes a classical Poisson-kernel / Schwarz-integral problem.", 0),
    ("Boundary values transform into a Fourier series in n; truncating at n≈200–1000 terms gives the sidewall flux Φ_swl (Case I) or Φ_swl+Φ_lat (Case II) to 6 significant figures.", 0),
    ("This flux directly gives the series coefficients A, B (Case I) or A′, B′ (Case II) that feed the threshold-voltage formula:", 0),
    ("Vth(W) = φ_ms + 2φ_f(1+a/W)/(1+b/W) + (Q′/Cox)/(1+b/W),   a=2tᵍ(B−A), b=−2tᵍA", 1),
], top=1.75, size=16.5)
footer(s)

# ---------------------------------------------------------------- 4. Verification story
s = add_slide()
title(s, "Verification: More Than Dimensional Analysis")
subtitle(s, "The paper's transcribed series (Eq. 11/17) did not survive independent scrutiny")
bullets(s, [
    ("Dimensional analysis alone caught two real transcription issues: a unit-inconsistent prefactor in A, and a missing ½ factor in sin(nktᵍ) — fixed.", 0),
    ("But fixing both was NOT sufficient: an independent finite-difference (FD) solve of the actual boundary-value problem (5-point Laplacian, sparse solve) still disagreed with the corrected series — even B, whose prefactor was never ambiguous, was off by ~2×.", 0),
    ("The FD solver itself was validated two ways, independent of Eq. 11/17:", 0),
    ("Sharp-step analytic limit (t_d,tᵍ→0): FD converges to a hand-derived closed form as the BVP simplifies (-0.4930 → -0.4659, vs. analytic -0.4652).", 1),
    ("Case I / Case II boundary continuity (t_d→t₁): two independently-coded BVPs agree in the shared limit — this check caught and fixed a real bug in the Case II lateral-wing flux integration bounds.", 1),
    ("Conclusion: FD solve is used as verified ground truth throughout; the root cause of the remaining series discrepancy is flagged as unresolved, not force-fit.", 0),
], top=1.75, size=15.5)
footer(s)

# ---------------------------------------------------------------- 5. Figure 5 result
s = add_slide()
title(s, "Reproduced Result: Case I / Case II Split (Fig. 5)")
_fh = 4.4
_fw = _fh * (825 / 675)
pic(s, f"{R}/li1990_fig5_repro.png", left=(13.333 - _fw) / 2, top=1.5, height=_fh)
bullets(s, [
    ("Vbg = 0, −1 V → Case I: ΔVth < 0 (inverse narrow-width effect, deepens as W shrinks)", 0),
    ("Vbg = −2, −4, −6 V → Case II: ΔVth > 0 (conventional effect)", 0),
    ("Case grouping matches the paper's own Fig. 5 legend exactly, bias-by-bias.", 0),
], top=6.15, size=13, height=1.3)
footer(s)

# ---------------------------------------------------------------- 6. Akers comparison
s = add_slide()
title(s, "Cross-Check Against Akers (1986)")
_fh2 = 4.4
_fw2 = _fh2 * (1650 / 675)
pic(s, f"{R}/li1990_vs_akers1986.png", left=(13.333 - _fw2) / 2, top=1.5, height=_fh2)
bullets(s, [
    ("Fully-recessed limit: FD-verified a=42.3 nm, b=71.8 nm vs. Akers' fringing length F=91 nm — same order of magnitude, independent derivations.", 0),
    ("Li et al. additionally predicts a bias-driven Case I → Case II sign reversal that Akers' single-regime model cannot represent.", 0),
], top=6.15, size=13, height=1.3)
footer(s)

# ---------------------------------------------------------------- 7. Takeaways
s = add_slide()
title(s, "Takeaways")
bullets(s, [
    ("Model reproduced end-to-end: 8 paper figures (Figs. 5–12), all matching qualitative trends and Case I/II groupings from the source figures.", 0),
    ("Rigor beyond pattern-matching: dimensional analysis found 2 bugs; an independent FD solve, cross-validated two separate ways, found the fixes were still incomplete and supplied a trustworthy ground truth.", 0),
    ("Consequence for numbers: fully-recessed length scales revised from the (flawed-series) a=103 nm/b=124 nm to the FD-verified a=42.3 nm/b=71.8 nm.", 0),
    ("Open item: the exact remaining error in the transcribed Eq. 11/17 series is not isolated — reported as unresolved rather than papered over.", 0),
    ("Full technical report: narrow_channel_semi_recessed_report.pdf (11 pp.) — this deck is a condensed companion.", 0),
], top=1.75, size=17)
footer(s)

prs.save(OUT)
print("Wrote", OUT)
